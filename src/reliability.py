"""Human reliability check.

  python src/reliability.py export --tag main     -> outputs/coding_sample_coder_A.xlsx and _B.xlsx
  python src/reliability.py score  --tag main     -> outputs/tables/reliability_*.csv

The sample is a stratified random 20% of valid responses (strata: platform x category; seed fixed).
Each coder receives an identical workbook WITHOUT platform names or the LLM's codes. The two
humans code independently. Completed workbooks go to data/manual/coding_sample_coder_A.xlsx and
data/manual/coding_sample_coder_B.xlsx; a consensus file (data/manual/coding_consensus.xlsx) is
made after discussing disagreements. This script never fills in human codes.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import CONFIG, INTERIM, MANUAL, ROOT, TABLES, brands_cfg, platforms_cfg  # noqa: E402

FRACTION = 0.20
STANCE = ["recommended", "mentioned_neutral", "discouraged", "not_present"]
SENT = ["positive", "neutral", "mixed", "negative", "not_present"]
FORM = ["list", "single", "comparison", "no_recommendation"]


def draw_sample(tag):
    r = pd.read_csv(INTERIM / f"responses_{tag}.csv", keep_default_na=False, dtype={"text": str})
    r = r[r["valid"].astype(str).str.lower() == "true"]
    seed = platforms_cfg()["seed"] + 11
    parts = []
    for _, g in r.groupby(["platform", "category"]):
        parts.append(g.sample(n=max(1, round(len(g) * FRACTION)), random_state=seed))
    s = pd.concat(parts).sample(frac=1, random_state=seed)  # shuffle so platforms are interleaved
    s["sample_no"] = range(1, len(s) + 1)
    s[["sample_no", "response_id", "platform", "category", "prompt_id"]].to_csv(MANUAL / "coding_sample_key.csv", index=False)
    return s


def export(tag):
    MANUAL.mkdir(parents=True, exist_ok=True)
    s = draw_sample(tag)
    prompts = pd.read_csv(CONFIG / "prompts.csv", dtype=str).set_index("prompt_id")
    codes = pd.read_csv(CONFIG / "attribute_codes.csv")
    brands = brands_cfg()["brands"]
    for coder in ("A", "B"):
        wb = Workbook()
        ins = wb.active; ins.title = "Instructions"
        lines = [
            f"Reliability coding: coder {coder}. Code independently; do not discuss with the other coder until both are done.",
            "Read config/codebook.md (the coding manual) before starting. It is the same manual the LLM coder used.",
            "Sheet 'Responses': one row per answer. Fill answer_form and top_pick (brand name or none).",
            "Sheet 'Brands': one row per answer x study brand. Fill is_brand_reference, stance, sentiment and attributes.",
            "attributes: attribute codes separated by semicolons (see sheet 'Codes'), or leave blank if none.",
            "If is_brand_reference is no, set stance and sentiment to not_present and leave attributes blank.",
            "Do not look up the platform; it is hidden on purpose. Do not edit the grey columns.",
            f"Sample: {len(s)} responses, stratified random 20% by platform and category (fixed seed).",
        ]
        for i, l in enumerate(lines, 1):
            ins.cell(i, 1, l)
        ins.column_dimensions["A"].width = 120
        grey = PatternFill("solid", fgColor="EEEEEE")

        ws = wb.create_sheet("Responses")
        head = ["sample_no", "question", "answer", "answer_form", "top_pick", "truncated", "notes"]
        ws.append(head)
        for x in s.itertuples():
            ws.append([x.sample_no, prompts.loc[x.prompt_id, "text"], x.text, "", "", "", ""])
        dv = DataValidation(type="list", formula1='"' + ",".join(FORM) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"D2:D{len(s) + 1}")
        dvt = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
        ws.add_data_validation(dvt); dvt.add(f"F2:F{len(s) + 1}")
        for col, w in zip("ABCDEFG", [9, 40, 100, 18, 18, 10, 30]):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row[:3]:
                c.fill = grey; c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in ws[1]:
            c.font = Font(bold=True)

        wsb = wb.create_sheet("Brands")
        wsb.append(["sample_no", "brand", "is_brand_reference", "stance", "sentiment", "attributes", "notes"])
        n = 1
        for x in s.itertuples():
            for b in brands:
                if b["category"] == x.category:
                    wsb.append([x.sample_no, b["name"], "", "", "", "", ""]); n += 1
        for col, vals in (("C", ["yes", "no"]), ("D", STANCE), ("E", SENT)):
            d = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
            wsb.add_data_validation(d); d.add(f"{col}2:{col}{n}")
        for col, w in zip("ABCDEFG", [9, 14, 18, 20, 14, 40, 30]):
            wsb.column_dimensions[col].width = w
        for row in wsb.iter_rows(min_row=2):
            for c in row[:2]:
                c.fill = grey
        for c in wsb[1]:
            c.font = Font(bold=True)

        wc = wb.create_sheet("Codes")
        wc.append(["category", "code", "label", "definition"])
        for r in codes.itertuples():
            wc.append([r.category, r.code, r.label, r.definition])
        wc.column_dimensions["D"].width = 100
        out = ROOT / "outputs" / f"coding_sample_coder_{coder}.xlsx"
        wb.save(out)
        print("wrote", out)


def kappa(a, b):
    a, b = pd.Series(a).astype(str), pd.Series(b).astype(str)
    cats = sorted(set(a) | set(b))
    if len(cats) == 1:
        return 1.0
    po = (a.values == b.values).mean()
    pe = sum((a == c).mean() * (b == c).mean() for c in cats)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


def landis_koch(k):
    if k is None or np.isnan(k):
        return ""
    return ("poor" if k < 0 else "slight" if k <= .20 else "fair" if k <= .40 else "moderate" if k <= .60
            else "substantial" if k <= .80 else "almost perfect")


def read_human(path):
    wb = load_workbook(path, data_only=True)
    r = pd.DataFrame(wb["Responses"].values); r.columns = r.iloc[0]; r = r.iloc[1:]
    b = pd.DataFrame(wb["Brands"].values); b.columns = b.iloc[0]; b = b.iloc[1:]
    return r, b


def score(tag):
    key = pd.read_csv(MANUAL / "coding_sample_key.csv")
    files = {c: MANUAL / f"coding_sample_coder_{c}.xlsx" for c in ("A", "B")}
    if not all(f.exists() for f in files.values()):
        sys.exit("Completed coder workbooks not found in data/manual/. Human coding is still pending.")
    (ra, ba), (rb, bb) = read_human(files["A"]), read_human(files["B"])
    name2id = {b["name"]: b["id"] for b in brands_cfg()["brands"]}
    rows = []
    bm = ba.merge(bb, on=["sample_no", "brand"], suffixes=("_A", "_B"))
    for v in ["is_brand_reference", "stance", "sentiment"]:
        k = kappa(bm[f"{v}_A"], bm[f"{v}_B"])
        rows.append({"variable": v, "level": "brand-in-response", "n": len(bm), "agreement": (bm[f"{v}_A"].astype(str) == bm[f"{v}_B"].astype(str)).mean(),
                     "kappa": k, "interpretation": landis_koch(k), "comparison": "human A vs human B"})
    codes = pd.read_csv(CONFIG / "attribute_codes.csv").code.unique()
    pres = bm[(bm.is_brand_reference_A == "yes") & (bm.is_brand_reference_B == "yes")]
    ks = []
    for c in codes:
        ia = pres.attributes_A.fillna("").map(lambda s: c in [x.strip() for x in str(s).split(";")])
        ib = pres.attributes_B.fillna("").map(lambda s: c in [x.strip() for x in str(s).split(";")])
        if ia.any() or ib.any():
            ks.append(kappa(ia, ib))
    rows.append({"variable": "attributes (mean kappa over codes)", "level": "brand-in-response", "n": len(pres),
                 "agreement": np.nan, "kappa": float(np.mean(ks)) if ks else np.nan,
                 "interpretation": landis_koch(np.mean(ks)) if ks else "", "comparison": "human A vs human B"})
    rm = ra.merge(rb, on="sample_no", suffixes=("_A", "_B"))
    for v in ["answer_form", "top_pick"]:
        k = kappa(rm[f"{v}_A"].astype(str).str.strip().str.lower(), rm[f"{v}_B"].astype(str).str.strip().str.lower())
        rows.append({"variable": v, "level": "response", "n": len(rm), "agreement": np.nan, "kappa": k,
                     "interpretation": landis_koch(k), "comparison": "human A vs human B"})
    cons_path = MANUAL / "coding_consensus.xlsx"
    if cons_path.exists():
        rc, bc = read_human(cons_path)
        llm = pd.read_csv(INTERIM / f"coded_brands_{tag}.csv")
        llm = llm[llm["pass"] == 1]
        bc = bc.merge(key, on="sample_no")
        bc["brand_id"] = bc.brand.map(name2id)
        j = bc.merge(llm, on=["response_id", "brand_id"], suffixes=("_h", "_llm"))
        j["ref_h"] = j.is_brand_reference.eq("yes")
        for v, h, l in [("is_brand_reference", "ref_h", "llm_reference"), ("stance", "stance_h", "stance_llm"),
                        ("sentiment", "sentiment_h", "sentiment_llm")]:
            k = kappa(j[h].astype(str).str.lower(), j[l].astype(str).str.lower())
            rows.append({"variable": v, "level": "brand-in-response", "n": len(j), "agreement": np.nan, "kappa": k,
                         "interpretation": landis_koch(k), "comparison": "LLM vs human consensus"})
    # LLM test-retest (pass 1 vs pass 2)
    llm = pd.read_csv(INTERIM / f"coded_brands_{tag}.csv")
    p1, p2 = llm[llm["pass"] == 1], llm[llm["pass"] == 2]
    tt = p1.merge(p2, on=["response_id", "brand_id"], suffixes=("_1", "_2"))
    for v in ["llm_reference", "stance", "sentiment"]:
        if len(tt):
            k = kappa(tt[f"{v}_1"], tt[f"{v}_2"])
            rows.append({"variable": v, "level": "brand-in-response", "n": len(tt), "agreement": (tt[f"{v}_1"].astype(str) == tt[f"{v}_2"].astype(str)).mean(),
                         "kappa": k, "interpretation": landis_koch(k), "comparison": "LLM run 1 vs run 2"})
    out = pd.DataFrame(rows)
    out.to_csv(TABLES / "reliability.csv", index=False)
    pd.DataFrame([("n_double_coded", len(key))], columns=["item", "value"]).to_csv(TABLES / "reliability_summary.csv", index=False)
    print(out.round(3).to_string())
    low = out[(out.comparison == "human A vs human B") & (out.kappa < 0.60)]
    if len(low):
        print("\nFLAG: kappa below 0.60 for:", ", ".join(low.variable))


def retest_only(tag):
    """LLM test-retest agreement, available before human coding is done."""
    llm = pd.read_csv(INTERIM / f"coded_brands_{tag}.csv")
    p1, p2 = llm[llm["pass"] == 1], llm[llm["pass"] == 2]
    tt = p1.merge(p2, on=["response_id", "brand_id"], suffixes=("_1", "_2"))
    rows = []
    for v in ["llm_reference", "stance", "sentiment"]:
        k = kappa(tt[f"{v}_1"], tt[f"{v}_2"])
        rows.append({"variable": v, "n": len(tt), "agreement": (tt[f"{v}_1"].astype(str) == tt[f"{v}_2"].astype(str)).mean(),
                     "kappa": k, "interpretation": landis_koch(k)})
    a1 = tt.attributes_1.fillna("").map(lambda s: set(str(s).split("|")) - {""})
    a2 = tt.attributes_2.fillna("").map(lambda s: set(str(s).split("|")) - {""})
    jac = [len(x & y) / len(x | y) if (x | y) else 1.0 for x, y in zip(a1, a2)]
    rows.append({"variable": "attributes (mean Jaccard)", "n": len(tt), "agreement": float(np.mean(jac)), "kappa": np.nan, "interpretation": ""})
    out = pd.DataFrame(rows)
    out.to_csv(TABLES / "llm_test_retest.csv", index=False)
    print(out.round(3).to_string())


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["export", "score", "retest"])
    ap.add_argument("--tag", default="main")
    a = ap.parse_args()
    {"export": export, "score": score, "retest": retest_only}[a.cmd](a.tag)
